#!/usr/bin/env python3
"""Writes the blank onboarding templates: docs/templates/*.csv and docs/templates/new_make_template.xlsx,
plus the pre-filled docs/templates/sinotruk_c7h_template.xlsx (profile + vehicle rows, ready for parts).

Each sheet has its header row and one example row starting with "#" — onboard.py skips rows whose first
cell starts with "#", so a template can be filled in place. The workbook adds a "help" tab and a note
on every column header.  Run:  python tools/onboard/make_templates.py
"""
import csv
from pathlib import Path

import sys
sys.path.insert(0, str(Path(__file__).parent))
from onboard import SHEETS  # noqa: E402

OUT = Path(__file__).resolve().parents[2] / "docs" / "templates"

HELP = {   # column -> what to put in it
    "code": "Short code, CAPITALS/digits/_ (profile: SINOTRUK_C7H; section: BMP; supplier: SUP-ACME)",
    "make": "Vehicle make exactly as it should show in CTP and Zoho Manufacturer (e.g. Sinotruk). Universal parts: 'All makes'",
    "model": "Vehicle model (= Zoho Brand), e.g. C7H. Blank for universal / all-makes parts",
    "display_name": "Name people see, e.g. 'Sinotruk C7H'",
    "parts_brand_code": "Who MAKES the parts: brand code (FAW, MANN, ...)", "parts_brand_name": "Parts maker's name",
    "parts_brand_is_oem": "Y if the parts are the vehicle maker's own (OEM)",
    "zoho_manufacturer": "Must equal 'make' (Zoho copies it back into CTP)", "zoho_brand": "Must equal 'model'",
    "zoho_sales_account_id": "Zoho GL account id for sales, e.g. 'Sales - Sinotruk' (onboard.py accounts lists them)",
    "zoho_purchase_account_id": "Zoho GL id for cost of sales, e.g. 'Cost of Goods Sold - Sinotruk'",
    "zoho_inventory_account_id": "Zoho GL id for stock: 'Inventory Asset' = 1429405000000034001",
    "unit": "Zoho unit, normally pcs", "default_part_type": "OEM, Aftermarket or Universal (used when a part row leaves it blank)",
    "rebate_scheme": "CTP rebate scheme / Zoho price book name the new parts join (optional)",
    "pn_pattern": "Regex a part number should match (warning only), e.g. ^[0-9]{4,}", "lookup_url": "Supplier lookup, {pn} = part number",
    "description_template": "Zoho item description, {make} {model} {sku}", "is_test": "Y only for a test brand (allows purge)",
    "notes": "Free text", "variant": "Cab / wheelbase / version, blank = the plain model", "year_from": "First model year (number)",
    "year_to": "Last model year (blank = still made)", "engine": "Engine code if it matters for fitment",
    "name": "Name as customers see it", "sku": "Part number = CTP SKU = Zoho SKU. No spaces. Must not exist for another make",
    "section_code": "Section (category) code: an existing one (BMP, LGT, MRR, ...) or one from sections",
    "part_type": "OEM, Aftermarket or Universal", "universal": "Y = fits every make (mechanical for all); needs no fitment rows",
    "catalogue_pn": "Number as printed in the maker's catalogue", "mpn": "Manufacturer part number if different",
    "side": "L, R, C (centre) or B (both); blank if none", "drawing_no": "Catalogue drawing number", "diagram_item_no": "Callout number on the drawing",
    "description": "Longer description", "uom": "EA unless otherwise", "weight_g": "Weight in grams",
    "list_price_zar": "Selling price excl. VAT in rands. Only used for NEW parts (Zoho stays the price master)",
    "supplier": "Preferred supplier's name (text shown on the part)", "note": "Free text",
    "xref_number": "The other number", "xref_brand": "Whose number it is", "xref_type": "oem, aftermarket, competitor, supersession or alternate",
    "confidence": "0-100 (100 = certain)", "supplier_code": "Code from suppliers (or an existing CTP supplier)", "supplier_sku": "Supplier's own code",
    "cost_zar": "Cost excl. VAT in rands (preferred supplier's cost becomes the new Zoho item's purchase rate)",
    "currency": "ZAR unless otherwise", "lead_time_days": "Days", "moq": "Minimum order quantity", "preferred": "Y for ONE supplier per part",
    "drawing_key": "Unique drawing code, e.g. SINOTRUK_C7H-BMP (never SEC<number>: those are FAW JH6's)",
    "title": "Drawing title", "image_file": "Picture file, relative to this sheet's folder (pictures/...)",
    "img_w": "Leave blank: read from the picture", "img_h": "Leave blank: read from the picture",
    "item_no": "Callout number on the drawing", "x": "Pixels from the LEFT of the drawing picture", "y": "Pixels from the TOP",
    "radius": "Leave blank (default marker size, about 3% of the drawing width)", "is_primary": "Y for the main photo (first photo if none)",
    "sort_order": "0, 1, 2 ...", "caption": "Optional", "scheme": "Rebate scheme name (must exist in CTP)",
    "contact": "Contact person", "phone": "Phone", "email": "Email", "vat_no": "VAT number",
}
EXAMPLE = {
    "profile": {"code": "SINOTRUK_C7H", "make": "Sinotruk", "model": "C7H", "display_name": "Sinotruk C7H", "parts_brand_code": "SINOTRUK",
                "parts_brand_name": "Sinotruk Genuine", "parts_brand_is_oem": "Y", "zoho_manufacturer": "Sinotruk", "zoho_brand": "C7H",
                "zoho_inventory_account_id": "1429405000000034001", "unit": "pcs", "default_part_type": "OEM"},
    "vehicles": {"model": "C7H", "year_from": "2018"}, "sections": {"code": "FLT", "name": "Filters"},
    "suppliers": {"code": "SUP-ACME", "name": "Acme Truck Parts", "currency": "ZAR", "lead_time_days": "21"},
    "parts": {"sku": "WG9725000001", "name": "Front Bumper Centre", "section_code": "BMP", "part_type": "OEM", "side": "C", "list_price_zar": "4250.00"},
    "fitment": {"sku": "WG9725000001", "model": "C7H", "year_from": "2018"},
    "xrefs": {"sku": "WG9725000001", "xref_number": "AB-1234", "xref_brand": "Acme", "xref_type": "aftermarket", "confidence": "90"},
    "part_suppliers": {"sku": "WG9725000001", "supplier_code": "SUP-ACME", "supplier_sku": "AC-778", "cost_zar": "1900.00", "preferred": "Y"},
    "diagrams": {"drawing_key": "SINOTRUK_C7H-BMP", "section_code": "BMP", "title": "Front Bumper — overview", "image_file": "pictures/bumper.png"},
    "hotspots": {"drawing_key": "SINOTRUK_C7H-BMP", "item_no": "1", "sku": "WG9725000001", "x": "812", "y": "440"},
    "images": {"sku": "WG9725000001", "image_file": "pictures/WG9725000001.jpg", "is_primary": "Y"},
    "price_books": {"scheme": "FAW JH6 – Rebate", "sku": "WG9725000001"},
}
# Real rows (not examples) for the pre-filled Sinotruk C7H workbook. GL account ids stay empty until Michelle
# confirms "Sales - Sinotruk" / "Cost of Goods Sold - Sinotruk" (or picks existing ones): validate warns and the
# Zoho step refuses until they are filled. pn_pattern / lookup_url / rebate_scheme are Ian's to fill.
PREFILL = {
    "sinotruk_c7h_template.xlsx": {
        "profile": [{"code": "SINOTRUK_C7H", "make": "Sinotruk", "model": "C7H", "display_name": "Sinotruk C7H",
                     "parts_brand_code": "SINOTRUK", "parts_brand_name": "Sinotruk Genuine", "parts_brand_is_oem": "Y",
                     "zoho_manufacturer": "Sinotruk", "zoho_brand": "C7H", "zoho_sales_account_id": "",
                     "zoho_purchase_account_id": "", "zoho_inventory_account_id": "1429405000000034001", "unit": "pcs",
                     "default_part_type": "OEM", "description_template": "{make} {model} · part no. {sku}",
                     "notes": "PLACEHOLDER accounts, Michelle to confirm: sales = 'Sales - Sinotruk', "
                              "purchase = 'Cost of Goods Sold - Sinotruk' (put the Zoho account ids in the two empty columns)"}],
        "vehicles": [{"model": "C7H", "notes": "Add one row per variant (cab / wheelbase) with model years"}],
    },
}


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    for name, (req, opt) in SHEETS.items():
        cols = req + opt
        with (OUT / f"{name}.csv").open("w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f); w.writerow(cols)
            ex = EXAMPLE.get(name, {}); row = [ex.get(c, "") for c in cols]; row[0] = "# " + row[0]
            w.writerow(row)
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "help"
    ws.append(["CTP Core — new make / brand template. Playbook: docs/13_ADD_A_NEW_MAKE.md"])
    ws.append(["Bold yellow columns are required. Rows starting with # are examples and are skipped. Hover a header for what goes in it."])
    ws.append(["Fill one tab per sheet; leave unused tabs empty. Put pictures in a 'pictures' folder next to this file."])
    ws.column_dimensions["A"].width = 120
    for name, (req, opt) in SHEETS.items():
        s = wb.create_sheet(name); cols = req + opt; s.append(cols)
        for i, c in enumerate(cols, 1):
            cell = s.cell(row=1, column=i); cell.comment = Comment(HELP.get(c, ""), "CTP")
            if c in req: cell.font = Font(bold=True); cell.fill = PatternFill("solid", fgColor="FFF2A8")
            s.column_dimensions[cell.column_letter].width = max(12, len(c) + 4)
        ex = EXAMPLE.get(name, {}); row = [ex.get(c, "") for c in cols]; row[0] = "# " + row[0]; s.append(row)
    wb.save(OUT / "new_make_template.xlsx")
    for fname, rows in PREFILL.items():
        wb = openpyxl.load_workbook(OUT / "new_make_template.xlsx")
        wb["help"].cell(row=1, column=1).value += f"  —  pre-filled: {fname} (profile + vehicles done; add parts etc.)"
        for name, recs in rows.items():
            s = wb[name]; cols = [c.value for c in s[1]]; s.delete_rows(2)   # drop the example row
            for r in recs: s.append([r.get(c, "") for c in cols])
        wb.save(OUT / fname)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
