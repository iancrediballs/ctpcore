"""CTP Core — Phase 3 stock reconciliation (read-only), 23 Oct 2026.

Inputs (pulled read-only; see docs/08_PHASE3_STOCK_TO_ZOHO.md):
  zoho_items.json    zoho-auth {action:"stock_report", part:"items"}
  zoho_details.json  ... part:"details"   (committed / available-for-sale / opening)
  zoho_docs.json     ... part:"docs"      (bills, invoices, sales orders, credit notes + lines)
  ctp_ledger.json    CTP stock_movement summed per part and reason (supabase db query -o json)

Per item:  difference = Zoho on hand − CTP on hand (CTP on hand excludes holds),
split exactly into
  stock_in     Zoho stock in (implied opening + approved bills) − CTP stock in (the 28 Jun opening
               receipt). Zoho's shelf arrived mostly on the Supalift bills PI000001–004 (19 Oct);
               CTP's on its 28 Jun opening load — the same goods, booked differently.
  sales_gap    CTP sales − Zoho net invoiced (positive = CTP took out more than Zoho invoiced)
  ctp_only     CTP adjustments + returns + transfers (none of these exist in Zoho)
so that difference = stock_in + sales_gap − ctp_only.

Usage: python tools/stock_reconcile.py <dir with the json files> <out prefix>
"""
import csv, json, sys
from collections import defaultdict

src, out = sys.argv[1], sys.argv[2]
load = lambda f: json.load(open(f"{src}/{f}", encoding="utf-8-sig"))
items = {i["item_id"]: i for i in load("zoho_items.json")}
details = {d["item_id"]: d for d in load("zoho_details.json")}
docs = load("zoho_docs.json")["docs"]
ctp = load("ctp_ledger.json")["rows"]

APPROVED_BILL = {"open", "overdue", "paid", "partially_paid"}
APPROVED_INV = {"sent", "paid", "partially_paid", "overdue"}
bills_in, inv_out, cn_back = defaultdict(float), defaultdict(float), defaultdict(float)
bill_refs, inv_refs = defaultdict(list), defaultdict(list)
per_bill = defaultdict(lambda: defaultdict(float))   # item -> bill number -> qty
for d in docs:
    for l in d.get("lines", []):
        k = l["item_id"]
        if d["kind"] == "bill" and d["status"] in APPROVED_BILL:
            bills_in[k] += l["qty"]; bill_refs[k].append(d["number"]); per_bill[k][d["number"]] += l["qty"]
        elif d["kind"] == "invoice" and d["status"] in APPROVED_INV:
            inv_out[k] += l["qty"]; inv_refs[k].append(d["number"])
        elif d["kind"] == "creditnote" and d["status"] in {"open", "closed"}:
            cn_back[k] += l["qty"]

n = lambda x: float(x or 0)
# A part linked to two Zoho items (an old INACTIVE one and the live one) appears
# twice in the ledger query; keep the active link only, or its stock counts twice.
has_active = {c["part_id"] for c in ctp if c.get("zoho_status") == "active"}
ctp = [c for c in ctp if not (c.get("zoho_status") == "inactive" and c["part_id"] in has_active)]
rows, seen = [], set()
for c in ctp:
    zid = c.get("zoho_item_id")
    if c["deleted"] and n(c["ctp_on_hand"]) == 0 and not zid:
        continue
    r = c["by_reason"] or {}
    z = items.get(zid) if zid else None
    dt = details.get(zid, {}) if zid else {}
    if zid: seen.add(zid)
    ctp_open = n(r.get("receipt"))
    ctp_sold = -n(r.get("sale"))
    ctp_only = n(r.get("adjustment")) + n(r.get("return")) + n(r.get("transfer"))
    ctp_on = n(c["ctp_on_hand"])
    row = {"sku": c["sku"], "name": c["name"], "ctp_part_id": c["part_id"], "zoho_item_id": zid or "",
           "zoho_tracked": bool(z and z.get("item_type") == "inventory"),
           "ctp_on_hand": ctp_on, "ctp_held": n(c["ctp_held"]), "ctp_available": n(c["ctp_available"]),
           "ctp_opening": ctp_open, "ctp_sold": ctp_sold, "ctp_adjust_return": ctp_only}
    if z and z.get("item_type") == "inventory":
        zon = n(dt.get("stock_on_hand", z.get("stock_on_hand")))
        b, o = bills_in.get(zid, 0), inv_out.get(zid, 0) - cn_back.get(zid, 0)
        zopen = zon - b + o
        row.update({"zoho_on_hand": zon, "zoho_committed": n(dt.get("committed_stock")),
                    "zoho_available_for_sale": n(dt.get("available_for_sale_stock", z.get("available_stock"))),
                    "zoho_opening_implied": zopen, "zoho_initial_stock_field": n(dt.get("initial_stock")),
                    "zoho_bills_in": b, "zoho_invoiced_net": o,
                    "zoho_bills": " ".join(sorted(set(bill_refs.get(zid, [])))), "zoho_invoices": " ".join(sorted(set(inv_refs.get(zid, []))))})
        diff = zon - ctp_on
        parts = {"stock_in": zopen + b - ctp_open, "sales_gap": ctp_sold - o, "ctp_only": -ctp_only}
        row.update({"difference": diff, "c_stock_in": parts["stock_in"],
                    "c_sales_gap": parts["sales_gap"], "c_ctp_only_adjust": parts["ctp_only"]})
        if abs(diff) < 1e-9:
            row["cause"] = "match"
            row["trust"] = "either (no action)"
        else:
            big = max(parts, key=lambda k: abs(parts[k]))
            causes = {
                "stock_in": f"stock-in differs: Zoho {zopen + b:g} (opening {zopen:g} + bills {b:g}{' ' + row['zoho_bills'] if row['zoho_bills'] else ''}) vs CTP opening {ctp_open:g}",
                "sales_gap": f"sales on one side only: CTP sold {ctp_sold:g}, Zoho invoiced net {o:g}",
                "ctp_only": f"CTP adjustments/returns {ctp_only:+g} never reached Zoho",
            }
            row["cause"] = "; ".join(causes[k] for k in sorted(parts, key=lambda k: -abs(parts[k])) if abs(parts[k]) > 1e-9)
            only = [k for k in parts if abs(parts[k]) > 1e-9]
            # The 28 Jun CTP opening load took ONE Supalift shipment's qty per SKU where the
            # item was on two bills: CTP opening == one bill's qty, and the gap == the other's.
            pb = per_bill.get(zid, {})
            one_bill = next(((kept, missed) for kept in pb for missed in pb if kept != missed
                             and abs(pb[kept] - ctp_open) < 1e-9 and abs(pb[missed] - parts["stock_in"]) < 1e-9 and zopen == 0), None)
            if one_bill:
                row["cause"] = (f"CTP opening load counted only {one_bill[0]} ({pb[one_bill[0]]:g}) and missed {one_bill[1]} "
                                f"({pb[one_bill[1]]:g}) — both Supalift shipments are on Zoho" +
                                ("; " + "; ".join(causes[k] for k in ("sales_gap", "ctp_only") if abs(parts[k]) > 1e-9) if len(only) > 1 else ""))
                row["trust"] = ("Zoho — spot-check a few to confirm the 2nd shipment is on the shelf" if only == ["stock_in"]
                                else "Zoho for the stock-in, then COUNT (other differences too)")
            elif only == ["ctp_only"]:
                row["trust"] = "CTP if the adjustment was a real count (then Zoho needs it) — confirm by counting"
            elif only == ["sales_gap"]:
                row["trust"] = "DOCUMENTS — find the missing invoice/sale; no count needed if found"
            elif "stock_in" not in only:
                row["trust"] = "DOCUMENTS + confirm by counting"
            else:
                row["trust"] = "COUNT IT (physical stock take)"
    else:
        row.update({"zoho_on_hand": "", "difference": "", "cause": "not linked to a tracked Zoho item" if not zid else "Zoho item not inventory-tracked",
                    "trust": "link / track in Zoho first" if ctp_on else "n/a (nothing on hand)"})
    rows.append(row)

for zid, z in items.items():
    if zid in seen or z.get("item_type") != "inventory":
        continue
    dt = details.get(zid, {})
    zon = n(dt.get("stock_on_hand", z.get("stock_on_hand")))
    rows.append({"sku": z.get("sku"), "name": z.get("name"), "ctp_part_id": "", "zoho_item_id": zid, "zoho_tracked": True,
                 "ctp_on_hand": "", "zoho_on_hand": zon, "zoho_committed": n(dt.get("committed_stock")),
                 "zoho_available_for_sale": n(dt.get("available_for_sale_stock")), "zoho_bills_in": bills_in.get(zid, 0),
                 "zoho_bills": " ".join(sorted(set(bill_refs.get(zid, [])))), "difference": "",
                 "cause": "in Zoho, not linked to any CTP part", "trust": "Zoho (or link it to a CTP part)" if zon else "n/a"})

cols = ["sku", "name", "ctp_part_id", "zoho_item_id", "zoho_tracked", "ctp_on_hand", "ctp_held", "ctp_available",
        "zoho_on_hand", "zoho_committed", "zoho_available_for_sale", "difference",
        "c_stock_in", "c_sales_gap", "c_ctp_only_adjust", "cause", "trust",
        "ctp_opening", "ctp_sold", "ctp_adjust_return", "zoho_opening_implied", "zoho_initial_stock_field",
        "zoho_bills_in", "zoho_invoiced_net", "zoho_bills", "zoho_invoices"]
rows.sort(key=lambda r: (-(abs(r["difference"]) if isinstance(r.get("difference"), float) else -1), str(r["sku"])))
with open(f"{out}.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore"); w.writeheader(); w.writerows(rows)

# totals
num = lambda k, sel=lambda r: True: sum(r[k] for r in rows if isinstance(r.get(k), (int, float)) and not isinstance(r.get(k), bool) and sel(r))
linked = lambda r: isinstance(r.get("difference"), float)
tot = {
    "items_in_report": len(rows),
    "linked_tracked_items": sum(1 for r in rows if linked(r)),
    "matching": sum(1 for r in rows if r.get("cause") == "match"),
    "differing": sum(1 for r in rows if linked(r) and r.get("cause") != "match"),
    "ctp_on_hand_total": num("ctp_on_hand"),
    "ctp_on_hand_linked": num("ctp_on_hand", linked),
    "ctp_held_total": num("ctp_held"),
    "zoho_on_hand_total": num("zoho_on_hand"),
    "zoho_on_hand_linked": num("zoho_on_hand", linked),
    "zoho_committed_total": num("zoho_committed"),
    "zoho_available_for_sale_total": num("zoho_available_for_sale"),
    "net_difference_linked": num("difference"),
    "abs_difference_linked": sum(abs(r["difference"]) for r in rows if linked(r)),
    "component_stock_in": num("c_stock_in"),
    "component_sales_gap": num("c_sales_gap"), "component_ctp_only_adjust": num("c_ctp_only_adjust"),
    "trust_zoho_spot_check": sum(1 for r in rows if str(r.get("trust", "")).startswith("Zoho — spot")),
    "zoho_stock_in_then_count": sum(1 for r in rows if str(r.get("trust", "")).startswith("Zoho for the stock-in")),
    "missed_shipment_units": sum(r["c_stock_in"] for r in rows if str(r.get("cause", "")).startswith("CTP opening load counted only")),
    "trust_ctp_confirm": sum(1 for r in rows if str(r.get("trust", "")).startswith("CTP")),
    "documents_only": sum(1 for r in rows if str(r.get("trust", "")).startswith("DOCUMENTS")),
    "count_it": sum(1 for r in rows if str(r.get("trust", "")).startswith("COUNT")),
    "zoho_only_items": sum(1 for r in rows if r.get("cause") == "in Zoho, not linked to any CTP part"),
    "zoho_only_units": num("zoho_on_hand", lambda r: r.get("cause") == "in Zoho, not linked to any CTP part"),
    "ctp_unlinked_with_stock": sum(1 for r in rows if not linked(r) and r.get("ctp_part_id") != "" and n(r.get("ctp_on_hand")) != 0),
}
json.dump(tot, open(f"{out}_totals.json", "w"), indent=1)

# xlsx: summary + full sheet
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
wb = Workbook(); ws = wb.active; ws.title = "Summary"
ws.append(["CTP Core stock reconciliation — 23 Oct 2026 (read-only; no adjustments made)"]); ws["A1"].font = Font(bold=True, size=13)
ws.append(["Zoho on hand vs CTP ledger on hand (CTP excludes holds). difference = Zoho − CTP = stock_in + sales_gap − ctp_only."])
ws.append([])
for k, v in tot.items(): ws.append([k.replace("_", " "), v])
full = wb.create_sheet("Per item"); full.append(cols)
for c in full[1]: c.font = Font(bold=True)
hl = {"COUNT": "FFF2CC", "DOCUMENTS": "DDEBF7", "CTP": "E2EFDA", "Zoho": "FCE4D6"}
for r in rows:
    full.append([r.get(k, "") for k in cols])
    t = str(r.get("trust", ""))
    for pre, colr in hl.items():
        if t.startswith(pre):
            for c in full[full.max_row]: c.fill = PatternFill("solid", fgColor=colr)
full.freeze_panes = "C2"; full.auto_filter.ref = full.dimensions
for col, wdt in zip("ABCDEFGHIJKLMNOPQR", [18, 34, 8, 20, 8, 8, 8, 8, 8, 8, 8, 9, 8, 8, 8, 8, 60, 34]):
    full.column_dimensions[col].width = wdt
docsheet = wb.create_sheet("Zoho documents")
docsheet.append(["kind", "number", "status", "date", "reference", "contact", "lines", "units"])
for d in docs:
    docsheet.append([d["kind"], d.get("number"), d.get("status"), d.get("date"), d.get("reference"), d.get("contact"),
                     len(d.get("lines", [])), sum(l["qty"] for l in d.get("lines", []))])
wb.save(f"{out}.xlsx")
print(json.dumps(tot, indent=1))
